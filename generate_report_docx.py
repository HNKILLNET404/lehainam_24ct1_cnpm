import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page setup: A4, Margins: Top 2cm, Bottom 2cm, Left 3cm, Right 2cm
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.79)     # 2cm
        section.bottom_margin = Inches(0.79)  # 2cm
        section.left_margin = Inches(1.18)    # 3cm
        section.right_margin = Inches(0.79)   # 2cm

    # Set default style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(13)
    font.color.rgb = RGBColor(0, 0, 0)

    # ================= TRANG BÌA =================
    p_b1 = doc.add_paragraph()
    p_b1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_b1.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\n")
    r1.font.bold = True
    r1.font.size = Pt(14)
    r2 = p_b1.add_run("TRƯỜNG ĐẠI HỌC KIẾN TRÚC ĐÀ NẴNG\n")
    r2.font.bold = True
    r2.font.size = Pt(15)
    r3 = p_b1.add_run("KHOA CÔNG NGHỆ THÔNG TIN\n")
    r3.font.bold = True
    r3.font.size = Pt(14)
    r_stars = p_b1.add_run("-------------------- *** --------------------\n\n\n")

    p_b2 = doc.add_paragraph()
    p_b2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p_b2.add_run("BÁO CÁO HỌC PHẦN\n")
    r4.font.bold = True
    r4.font.size = Pt(16)
    r5 = p_b2.add_run("MÔN: CÔNG NGHỆ PHẦN MỀM\n\n")
    r5.font.bold = True
    r5.font.size = Pt(18)
    r5.font.color.rgb = RGBColor(26, 82, 118)

    r6 = p_b2.add_run("ĐỀ TÀI:\n")
    r6.font.bold = True
    r6.font.size = Pt(16)
    r7 = p_b2.add_run("XÂY DỰNG HỆ THỐNG QUẢN LÝ THÔNG TIN\nWEBSITE BÁN ÁO THỜI TRANG STREETWEAR URBAN THREADS\n\n\n\n")
    r7.font.bold = True
    r7.font.size = Pt(18)
    r7.font.color.rgb = RGBColor(217, 83, 30)

    # Info table
    p_info = doc.add_paragraph()
    p_info.paragraph_format.left_indent = Inches(1.5)
    p_info.paragraph_format.line_spacing = 1.3
    
    r_i = p_info.add_run("Giảng viên hướng dẫn\t:  ThS. Phạm Thị Dung\n")
    r_i.font.bold = True
    r_i.font.size = Pt(13)

    r_i2 = p_info.add_run("Sinh viên thực hiện\t:  Lê Hải Nam\n")
    r_i2.font.bold = True
    r_i2.font.size = Pt(13)

    r_i3 = p_info.add_run("Mã số sinh viên (MSSV)\t:  2451220069\n")
    r_i3.font.bold = True
    r_i3.font.size = Pt(13)

    r_i4 = p_info.add_run("Lớp sinh hoạt\t\t:  24CT1\n")
    r_i4.font.bold = True
    r_i4.font.size = Pt(13)

    r_i5 = p_info.add_run("Khoa\t\t\t:  Công Nghệ Thông Tin\n\n\n\n")
    r_i5.font.bold = True
    r_i5.font.size = Pt(13)

    p_b3 = doc.add_paragraph()
    p_b3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_bot = p_b3.add_run("Đà Nẵng, Năm 2026")
    r_bot.font.italic = True
    r_bot.font.size = Pt(13)

    doc.add_page_break()

    # ================= LỜI NÓI ĐẦU =================
    h_intro = doc.add_heading("LỜI NÓI ĐẦU", level=1)
    h_intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in h_intro.runs:
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(26, 82, 118)

    p = doc.add_paragraph(
        "Trong thời đại chuyển đổi số và bùng nổ của thương mại điện tử toàn cầu, phương thức mua sắm của người tiêu dùng "
        "— đặc biệt là thế hệ trẻ — đã có những bước chuyển biến sâu sắc. Thay vì mua sắm trực tiếp truyền thống, người tiêu dùng "
        "ngày nay ưu tiên các nền tảng trực tuyến với khả năng tra cứu thông tin nhanh chóng, so sánh mẫu mã, lựa chọn biến thể kích cỡ "
        "và thanh toán linh hoạt chỉ qua vài thao tác đơn giản trên thiết bị thông minh. Đối với ngành công nghiệp thời trang đường phố (Streetwear), "
        "website không đơn thuần là một kênh giới thiệu sản phẩm mà còn là nơi thể hiện bản sắc thương hiệu, văn hóa lối sống và kết nối cộng đồng người tiêu dùng trẻ tuổi."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "Xuất phát từ nhu cầu thực tiễn đó, trong khuôn khổ học phần Công nghệ phần mềm dưới sự giảng dạy tận tình của cô Phạm Thị Dung, "
        "em đã lựa chọn và thực hiện đề tài: \"Xây dựng hệ thống quản lý thông tin Website bán áo thời trang Urban Threads\". "
        "Dự án tập trung nghiên cứu, thiết kế và phát triển một nền tảng thương mại điện tử hoàn chỉnh theo mô hình MVT (Model - View - Template) "
        "sử dụng ngôn ngữ Python và framework Django 5.x, kết hợp thiết kế giao diện hiện đại với Tailwind CSS và Alpine.js. "
        "Đặc biệt, hệ thống tích hợp giải pháp thanh toán chuyển khoản ngân hàng quét mã VietQR động hỗ trợ đa ngân hàng tại Việt Nam, "
        "cùng công cụ tìm kiếm tiếng Việt thông minh và hệ thống quản trị kiểm soát đơn hàng, tồn kho trực quan 1-click."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "Đề tài này là cơ hội quý báu giúp em củng cố kiến thức nền tảng về vòng đời phát triển phần mềm (SDLC), từ giai đoạn khảo sát yêu cầu, "
        "xây dựng mô hình toán học giải thuật, mô hình hóa Use Case, biểu đồ tuần tự cho đến lập trình thực nghiệm và kiểm thử. "
        "Em xin gửi lời cảm ơn chân thành đến cô Phạm Thị Dung đã truyền đạt kiến thức và đồng hành hướng dẫn em trong suốt quá trình hoàn thành báo cáo này."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    doc.add_page_break()

    # ================= MỤC LỤC =================
    h_toc = doc.add_heading("MỤC LỤC", level=1)
    h_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in h_toc.runs:
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(26, 82, 118)

    toc_items = [
        ("LỜI NÓI ĐẦU", 2),
        ("MỤC LỤC", 3),
        ("CHƯƠNG 1: TỔNG QUAN CƠ SỞ LÝ THUYẾT", 4),
        ("    1.1. Giới thiệu và cài đặt về công nghệ, ngôn ngữ, dữ liệu", 4),
        ("        1.1.1. Giới thiệu ngôn ngữ Python và Framework Django 5.x", 4),
        ("        1.1.2. Công nghệ xây dựng giao diện Tailwind CSS và Alpine.js", 5),
        ("        1.1.3. Hệ thống cơ sở dữ liệu và Cổng thanh toán VietQR", 6),
        ("        1.1.4. Hướng dẫn chi tiết cài đặt môi trường thực hiện đề tài", 7),
        ("    1.2. Phân tích yêu cầu hệ thống đề tài", 8),
        ("        1.2.1. Khảo sát các website tham khảo thực tế", 8),
        ("        1.2.2. Bảng phân tích yêu cầu chức năng (Functional Requirements)", 9),
        ("        1.2.3. Bảng phân tích yêu cầu phi chức năng (Non-Functional Requirements)", 11),
        ("        1.2.4. Mô hình toán học và giải thuật (Đăng ký & Đăng nhập)", 12),
        ("CHƯƠNG 2: BIỂU ĐỒ VÀ THIẾT KẾ HỆ THỐNG", 14),
        ("    2.1. Biểu đồ Use Case", 14),
        ("    2.2. Biểu đồ tuần tự (Sequence Diagram)", 16),
        ("    2.3. Sơ đồ cơ sở dữ liệu (ER Diagram)", 18),
        ("CHƯƠNG 3: TRÌNH BÀY DEMO HỆ THỐNG", 20),
        ("    3.1. Giao diện Trang chủ và Danh mục", 20),
        ("    3.2. Giao diện Đăng nhập, Đăng ký và Tìm kiếm thông minh", 21),
        ("    3.3. Trang chi tiết sản phẩm và chọn biến thể (Size / Màu)", 22),
        ("    3.4. Giỏ hàng và Ngăn kéo giỏ hàng (Mini Cart Drawer)", 23),
        ("    3.5. Trang thanh toán và Cổng quét mã VietQR tự động", 24),
        ("    3.6. Trang quản trị Admin Dashboard", 25),
        ("    3.7. Kế hoạch triển khai dự án (Gantt Chart)", 26),
        ("KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", 27),
        ("LỜI CẢM ƠN", 28),
    ]

    p_toc = doc.add_paragraph()
    p_toc.paragraph_format.line_spacing = 1.3
    for title_text, page_num in toc_items:
        r_t = p_toc.add_run(f"{title_text.ljust(65, '.')}{str(page_num).rjust(5)}\n")
        r_t.font.size = Pt(12)
        if "CHƯƠNG" in title_text or "KẾT LUẬN" in title_text or "LỜI" in title_text:
            r_t.font.bold = True

    doc.add_page_break()

    # ================= CHƯƠNG 1 =================
    h_c1 = doc.add_heading("CHƯƠNG 1: TỔNG QUAN CƠ SỞ LÝ THUYẾT", level=1)
    for r in h_c1.runs:
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(26, 82, 118)

    # 1.1
    h_11 = doc.add_heading("1.1. Giới thiệu và cài đặt về công nghệ, ngôn ngữ, dữ liệu", level=2)
    for r in h_11.runs:
        r.font.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(217, 83, 30)

    # 1.1.1
    h_111 = doc.add_heading("1.1.1. Giới thiệu ngôn ngữ Python và Framework Django 5.x", level=3)
    p = doc.add_paragraph(
        "Python là một trong những ngôn ngữ lập trình bậc cao phổ biến nhất hiện nay nhờ cú pháp rõ ràng, súc tích, "
        "thư viện tiêu chuẩn đồ sộ và khả năng áp dụng rộng rãi từ lập trình web, tự động hóa cho đến trí tuệ nhân tạo. "
        "Trên nền tảng Python, Django là framework phát triển web hàng đầu, được xây dựng theo triết lý \"Batteries-included\" "
        "— cung cấp sẵn hầu như toàn bộ các công cụ và linh kiện cần thiết để xây dựng một ứng dụng web hoàn chỉnh từ quy mô nhỏ đến cấp doanh nghiệp."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "Django hoạt động theo kiến trúc MVT (Model - View - Template), một biến thể hiện đại của mô hình MVC truyền thống:\n"
        "• Model (Mô hình dữ liệu): Chịu trách nhiệm định nghĩa cấu trúc dữ liệu và logic nghiệp vụ. Django ORM (Object-Relational Mapping) "
        "cho phép lập trình viên thao tác với bảng dữ liệu hoàn toàn bằng các đối tượng Python, không cần phải viết các câu lệnh SQL thủ công, "
        "giúp tăng tốc độ phát triển và ngăn chặn triệt để lỗ hổng SQL Injection.\n"
        "• View (Tầng điều khiển / Xử lý nghiệp vụ): Tiếp nhận yêu cầu HTTP từ người dùng thông qua URL dispatcher, tương tác với tầng Model "
        "để truy vấn hoặc cập nhật dữ liệu, xử lý nghiệp vụ bán hàng (thêm giỏ, áp coupon, sinh mã QR) và chuyển dữ liệu sang Template để hiển thị.\n"
        "• Template (Tầng trình bày giao diện): Sử dụng Django Template Language (DTL) để kết hợp mã HTML với các biến, bộ lọc định dạng tiền tệ (filters) "
        "và thẻ điều kiện (tags) nhằm kết xuất ra giao diện động gửi về cho trình duyệt của khách hàng."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "Các ưu điểm vượt trội của Django 5.x được khai thác trong đề tài:\n"
        "1. Hệ thống xác thực và phân quyền người dùng (Authentication System) tích hợp sẵn, hỗ trợ phân chia nhóm quyền quản trị (Superuser, Staff) và khách hàng.\n"
        "2. Giao diện quản trị tự động Django Admin cực kỳ mạnh mẽ, cho phép cấu hình tùy biến bộ lọc, tìm kiếm, hiển thị ảnh xem trước và thao tác hàng loạt.\n"
        "3. Cơ chế bảo mật đa lớp: Tự động gắn thẻ CSRF token chống giả mạo request, bảo vệ chống XSS (Cross-Site Scripting), Clickjacking và cơ chế băm mật khẩu chuẩn PBKDF2 với thuật toán SHA-256."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    # 1.1.2
    h_112 = doc.add_heading("1.1.2. Công nghệ xây dựng giao diện Tailwind CSS và Alpine.js", level=3)
    p = doc.add_paragraph(
        "Giao diện người dùng của website Urban Threads được thiết kế theo phong cách thời trang đường phố (Streetwear Monochrome) "
        "tối giản và hiện đại, sử dụng hai công nghệ frontend tiên tiến:\n"
        "• Tailwind CSS: Framework CSS theo xu hướng Utility-First, cung cấp hàng ngàn lớp tiện ích dựng sẵn giúp lập trình viên tùy biến giao diện trực tiếp "
        "ngay trong mã HTML mà không cần phải viết các tệp CSS riêng lẻ cồng kềnh. Tailwind CSS mang lại khả năng đáp ứng (Responsive Design) hoàn hảo trên "
        "tất cả kích thước màn hình từ điện thoại thông minh, máy tính bảng đến máy tính để bàn.\n"
        "• Alpine.js: Thư viện JavaScript siêu nhẹ (chỉ khoảng 15KB) nhưng mang đầy đủ tính năng phản ứng (Reactive) tương tự Vue.js hoặc React. "
        "Trong đề tài, Alpine.js được sử dụng để điều khiển các thành phần tương tác tức thì: Ngăn kéo giỏ hàng trượt ra từ bên phải màn hình (Mini Cart Drawer), "
        "bộ chuyển đổi biến thể sản phẩm (đổi màu, chọn size tự động đổi giá và kiểm tra tồn kho), bộ đếm thông báo Toast và xem phóng to ảnh sản phẩm."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    # 1.1.3
    h_113 = doc.add_heading("1.1.3. Hệ thống cơ sở dữ liệu và Cổng thanh toán VietQR", level=3)
    p = doc.add_paragraph(
        "• Cơ sở dữ liệu: Hệ thống sử dụng cơ sở dữ liệu SQLite trong môi trường phát triển và thử nghiệm nhờ ưu điểm nhỏ gọn, không cần cấu hình server riêng "
        "và tệp cơ sở dữ liệu được lưu trực tiếp tại db.sqlite3 trong thư mục gốc. Thông qua Django ORM, hệ thống hoàn toàn sẵn sàng chuyển đổi sang các hệ quản trị "
        "cơ sở dữ liệu mạnh mẽ như PostgreSQL hoặc MySQL khi triển khai thực tế trên máy chủ đám mây.\n"
        "• Cổng thanh toán VietQR động: Nhằm bắt kịp xu hướng thanh toán không dùng tiền mặt tại Việt Nam, đề tài xây dựng module tích hợp chuẩn VietQR (Napas247). "
        "Khi khách hàng đặt hàng, hệ thống tự động sinh mã QR chứa đúng số tài khoản thụ hưởng, số tiền chính xác của đơn hàng và cú pháp chuyển khoản DHxxxxxxxx. "
        "Hệ thống cho phép cấu hình linh hoạt sang bất kỳ ngân hàng nào tại Việt Nam (Vietcombank, Techcombank, MB Bank, ACB, BIDV, VPBank...) thông qua tệp môi trường .env "
        "mà không cần sửa đổi mã nguồn."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    # 1.1.4
    h_114 = doc.add_heading("1.1.4. Hướng dẫn chi tiết cài đặt môi trường thực hiện đề tài", level=3)
    p = doc.add_paragraph("Để triển khai và vận hành website Urban Threads trên môi trường cục bộ, các bước thực hiện như sau:")
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    steps = [
        ("Bước 1: Cài đặt Python", "Tải và cài đặt Python phiên bản 3.10 trở lên từ trang chủ python.org. Chú ý đánh dấu vào tùy chọn \"Add Python to PATH\" trước khi bấm Install."),
        ("Bước 2: Mở thư mục dự án", "Mở cửa sổ dòng lệnh (Command Prompt hoặc PowerShell) và điều hướng đến thư mục dự án: cd D:\\urban-threads"),
        ("Bước 3: Tạo môi trường ảo (Virtual Environment)", "Tạo môi trường ảo độc lập để cô lập các gói thư viện: python -m venv venv"),
        ("Bước 4: Kích hoạt môi trường ảo", "Kích hoạt môi trường ảo: venv\\Scripts\\activate (Trên thanh dòng lệnh sẽ xuất hiện tiền tố (venv))."),
        ("Bước 5: Cài đặt các thư viện phụ thuộc", "Cài đặt toàn bộ các thư viện được định nghĩa trong tệp requirements.txt bằng lệnh: pip install -r requirements.txt (gồm Django, Pillow, qrcode, python-decouple, django-crispy-forms, crispy-tailwind...)."),
        ("Bước 6: Khởi chạy máy chủ phát triển", "Chạy lệnh: python manage.py runserver. Truy cập website tại địa chỉ http://127.0.0.1:8000/ và trang quản trị tại http://127.0.0.1:8000/admin/.")
    ]

    for title_s, desc_s in steps:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.left_indent = Inches(0.4)
        p_s.paragraph_format.line_spacing = 1.3
        r_t = p_s.add_run(f"• {title_s}: ")
        r_t.font.bold = True
        p_s.add_run(desc_s)

    # 1.2
    h_12 = doc.add_heading("1.2. Phân tích yêu cầu hệ thống đề tài", level=2)
    for r in h_12.runs:
        r.font.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor(217, 83, 30)

    # 1.2.1
    h_121 = doc.add_heading("1.2.1. Khảo sát các website tham khảo thực tế", level=3)
    p = doc.add_paragraph(
        "Để đảm bảo hệ thống vừa đáp ứng đầy đủ tính năng thương mại điện tử chuyên nghiệp, vừa mang phong cách thẩm mỹ thời thượng của ngành thời trang đường phố, "
        "dự án đã tiến hành khảo sát và chắt lọc kinh nghiệm từ 3 website thương mại điện tử hàng đầu:"
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    sites = [
        ("1. Chivalry Warehouse (https://chivalry-warehouse.com/)", 
         "Local brand streetwear nổi tiếng tại Việt Nam với phong cách Rock & Grunge độc đáo. Đề tài học hỏi thiết kế thanh điều hướng cố định (Sticky Header), "
         "cách phân loại danh mục sản phẩm thời trang (Áo thun, Hoodie, Quần, Phụ kiện, Sale), các nhãn dán trạng thái sản phẩm (New Arrival, Best Seller, Sale) "
         "và bộ chọn thuộc tính biến thể màu sắc (Color swatches) trực quan."),
        ("2. Represent Clo (https://representclo.com/)", 
         "Thương hiệu thời trang đường phố cao cấp hàng đầu Vương Quốc Anh (UK High-end Streetwear). Đề tài học hỏi phong cách thiết kế Monochrome tối giản sang trọng, "
         "bố cục lưới sản phẩm thông thoáng, thư viện ảnh gallery chi tiết nhiều góc chụp, bảng hướng dẫn chọn kích cỡ (Size Guide) và khu vực gợi ý sản phẩm liên quan (Related Products)."),
        ("3. Culture Kings (https://culturekings.com.au/)", 
         "Chuỗi bán lẻ streetwear và sneakers quy mô lớn nhất Australia. Đề tài học hỏi và triển khai tính năng Ngăn kéo giỏ hàng trượt ra từ bên phải màn hình (Mini Cart Drawer) "
         "mà không cần tải lại trang, thanh tiến trình miễn phí vận chuyển (Free shipping progress bar đạt mốc 500.000đ), cơ chế lọc sản phẩm đa tiêu chí và bộ tìm kiếm thông minh.")
    ]

    for site_name, site_desc in sites:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Inches(0.4)
        p_st.paragraph_format.line_spacing = 1.3
        r_sn = p_st.add_run(f"{site_name}\n")
        r_sn.font.bold = True
        r_sn.font.color.rgb = RGBColor(26, 82, 118)
        p_st.add_run(site_desc)

    # 1.2.2
    h_122 = doc.add_heading("1.2.2. Bảng phân tích yêu cầu chức năng (Functional Requirements - FR)", level=3)
    p = doc.add_paragraph("Các yêu cầu chức năng của hệ thống Urban Threads được phân chia và quản lý chi tiết qua bảng tiêu chuẩn dưới đây:")
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    # Table FR
    table_fr = doc.add_table(rows=1, cols=7)
    table_fr.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_fr.autofit = False

    fr_headers = ["Mã", "Tên yêu cầu", "Mô tả yêu cầu", "Ưu tiên", "Tiêu chí nghiệm thu", "Tình trạng", "Link mẫu"]
    hdr_cells = table_fr.rows[0].cells
    for i, h_text in enumerate(fr_headers):
        hdr_cells[i].text = h_text
        set_cell_background(hdr_cells[i], "D9531E")
        set_cell_margins(hdr_cells[i])
        p_h = hdr_cells[i].paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p_h.runs:
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)

    fr_data = [
        ("FR-01", "Đăng ký tài khoản", "Người dùng có thể đăng ký tài khoản mới bằng Email và mật khẩu bảo mật.", "Cao", "Nhập đúng thông tin tạo tài khoản thành công; trùng email hoặc sai định dạng hiển thị thông báo lỗi.", "Hoàn thành", "chivalry-warehouse.com"),
        ("FR-02", "Đăng nhập / Đăng xuất", "Người dùng đăng nhập bằng Email/Password, có ghi nhớ phiên và đăng xuất an toàn.", "Cao", "Đăng nhập đúng vào trang cá nhân; sai mật khẩu báo lỗi; đăng xuất hủy session an toàn.", "Hoàn thành", "representclo.com"),
        ("FR-03", "Sổ địa chỉ giao hàng", "Người dùng quản lý nhiều địa chỉ nhận hàng (thêm, sửa, xóa, đặt địa chỉ mặc định).", "Trung bình", "Dữ liệu địa chỉ lưu vào tài khoản; tự động điền địa chỉ mặc định khi đặt hàng.", "Hoàn thành", "culturekings.com.au"),
        ("FR-04", "Lịch sử đơn hàng", "Người dùng theo dõi danh sách đơn hàng đã mua và tiến trình trạng thái chi tiết.", "Cao", "Hiển thị mã đơn DHxxxx, danh sách sản phẩm, giá tiền, ngày đặt và timeline trạng thái 7 bước.", "Hoàn thành", "representclo.com"),
        ("FR-05", "Duyệt danh mục & Bộ lọc", "Phân loại sản phẩm (Áo, Quần, Phụ kiện, Sale) kèm bộ lọc theo giá, màu, size.", "Cao", "Lọc chính xác theo nhiều tiêu chí đồng thời; sắp xếp theo giá tăng/giảm, mới nhất.", "Hoàn thành", "chivalry-warehouse.com"),
        ("FR-06", "Tìm kiếm thông minh", "Tìm kiếm sản phẩm theo từ khóa tiếng Việt đa tầng (hỗ trợ cả có dấu và không dấu).", "Cao", "Gõ \"ao thun\" hay \"áo thun\" đều trả về kết quả đúng; có popup gợi ý sản phẩm.", "Hoàn thành", "culturekings.com.au"),
        ("FR-07", "Chi tiết & Biến thể SKU", "Trang chi tiết sản phẩm hiển thị gallery ảnh, bảng chọn size/màu và tồn kho tương ứng.", "Cao", "Đổi màu/size tự động cập nhật mã SKU, giá bán và cảnh báo còn/hết hàng tức thì.", "Hoàn thành", "representclo.com"),
        ("FR-08", "Danh sách yêu thích", "Khách hàng có thể lưu các sản phẩm yêu thích (Wishlist) vào tài khoản cá nhân.", "Trung bình", "Bấm icon trái tim thêm/xóa nhanh bằng AJAX mà không cần tải lại trang.", "Hoàn thành", "chivalry-warehouse.com"),
        ("FR-09", "Giỏ hàng & Mini Cart", "Thêm sản phẩm vào giỏ, xem nhanh qua ngăn kéo trượt ra từ bên phải màn hình.", "Cao", "Mini Cart trượt mượt mà; tăng giảm số lượng, xóa item và cập nhật tổng tiền realtime.", "Hoành thành", "culturekings.com.au"),
        ("FR-10", "Mã giảm giá & Freeship", "Áp dụng coupon khuyến mãi và hiển thị thanh tiến trình đạt ngưỡng miễn phí ship.", "Cao", "Nhập mã hợp lệ tự trừ tiền; đơn hàng đạt từ 500,000đ tự động miễn phí vận chuyển.", "Hoàn thành", "culturekings.com.au"),
        ("FR-11", "Thanh toán VietQR động", "Tự động sinh mã VietQR động theo chuẩn Napas247 cho bất kỳ ngân hàng nào.", "Cao", "Mã QR chứa đúng số tiền, STK và nội dung DHxxxx; có nút 1-click copy nhanh.", "Hoàn thành", "VietQR Standard"),
        ("FR-12", "Upload biên lai đối soát", "Khách hàng tải lên ảnh chụp màn hình chuyển khoản sau khi quét mã QR.", "Cao", "Tải được ảnh JPG/PNG; tự động cập nhật trạng thái đơn sang \"Đã xác nhận\".", "Hoàn thành", "Urban Threads"),
        ("FR-13", "Quản lý sản phẩm & Kho", "Admin thêm, sửa, xóa sản phẩm, danh mục, gallery ảnh và quản lý tồn kho biến thể.", "Cao", "Dữ liệu lưu chuẩn CSDL; có thumbnail ảnh trực quan; kiểm soát tồn kho từng size/màu.", "Hoàn thành", "representclo.com"),
        ("FR-14", "Quản lý đơn hàng", "Theo dõi vòng đời đơn hàng qua 7 trạng thái, lọc theo thời gian và bulk update.", "Cao", "Cập nhật trạng thái chuẩn xác; tự động trừ tồn kho khi đơn hàng được xác nhận.", "Hoàn thành", "Django Admin"),
        ("FR-15", "Duyệt thanh toán 1-click", "Admin kiểm tra ảnh biên lai do khách gửi và duyệt thanh toán trực tiếp tại bảng đơn.", "Cao", "Xem trước ảnh bill phóng to; bấm 1 nút xác nhận tiền về và chuyển đơn sang đóng gói.", "Hoàn thành", "Urban Threads"),
        ("FR-16", "Quản lý khuyến mãi", "Tạo mã coupon (% hoặc số tiền), giới hạn số lượt dùng và quản lý banner trang chủ.", "Trung bình", "Mã hết hạn hoặc hết lượt sẽ tự động từ chối; banner cập nhật đúng thứ tự hiển thị.", "Hoàn thành", "culturekings.com.au"),
    ]

    for row_idx, r_data in enumerate(fr_data):
        row_cells = table_fr.add_row().cells
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F9F9F9"
        for c_idx, val in enumerate(r_data):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p_c = row_cells[c_idx].paragraphs[0]
            p_c.paragraph_format.line_spacing = 1.15
            if c_idx in [0, 3, 5]:
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p_c.runs:
                run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 1.2.3
    h_123 = doc.add_heading("1.2.3. Bảng phân tích yêu cầu phi chức năng (Non-Functional Requirements - NFR)", level=3)
    
    table_nfr = doc.add_table(rows=1, cols=7)
    table_nfr.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_nfr.autofit = False

    hdr_cells_n = table_nfr.rows[0].cells
    for i, h_text in enumerate(fr_headers):
        hdr_cells_n[i].text = h_text
        set_cell_background(hdr_cells_n[i], "1A5276")
        set_cell_margins(hdr_cells_n[i])
        p_h = hdr_cells_n[i].paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p_h.runs:
            run.font.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(255, 255, 255)

    nfr_data = [
        ("NFR-01", "Hiệu năng tải trang", "Trang chủ, danh mục và chi tiết sản phẩm phản hồi nhanh trong thời gian chấp nhận được.", "Cao", "Thời gian phản hồi trang dưới 2 giây với kết nối mạng thông thường; tối ưu truy vấn ORM.", "Hoàn thành", "Google Lighthouse"),
        ("NFR-02", "Bảo mật hệ thống", "Bảo vệ an toàn dữ liệu người dùng, mật khẩu và chống các lỗ hổng tấn công phổ biến.", "Cao", "Mã hóa mật khẩu chuẩn PBKDF2/SHA256, tự động kiểm tra CSRF token, chống XSS và SQL Injection.", "Hoàn thành", "Django Security Core"),
        ("NFR-03", "Thiết kế Responsive", "Giao diện hiển thị tối ưu trên mọi thiết bị: Điện thoại di động, Máy tính bảng và Laptop.", "Cao", "Layout co giãn linh hoạt bằng Tailwind CSS; menu mobile và mini cart drawer thao tác mượt mà.", "Hoàn thành", "representclo.com"),
        ("NFR-04", "Linh hoạt cấu hình QR", "Khả năng thay đổi cổng tài khoản ngân hàng dễ dàng mà không làm ảnh hưởng đến mã nguồn.", "Cao", "Chỉnh sửa BANK_ID trong tệp .env (VCB, TCB, MB, ACB, BIDV...) hệ thống tự nhận diện ngân hàng tương ứng.", "Hoàn thành", "VietQR Standard"),
    ]

    for row_idx, r_data in enumerate(nfr_data):
        row_cells = table_nfr.add_row().cells
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F9F9F9"
        for c_idx, val in enumerate(r_data):
            row_cells[c_idx].text = val
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p_c = row_cells[c_idx].paragraphs[0]
            p_c.paragraph_format.line_spacing = 1.15
            if c_idx in [0, 3, 5]:
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p_c.runs:
                run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 1.2.4
    h_124 = doc.add_heading("1.2.4. Mô hình toán học và giải thuật (Đăng ký & Đăng nhập)", level=3)
    p = doc.add_paragraph(
        "Theo chuẩn phân tích yêu cầu công nghệ phần mềm, mô hình giải thuật toán học của hai chức năng nền tảng "
        "— Đăng ký tài khoản và Đăng nhập hệ thống — được mô hình hóa theo không gian trạng thái toán học và lưu đồ giải thuật như sau:"
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p_m1 = doc.add_paragraph()
    p_m1.paragraph_format.left_indent = Inches(0.4)
    p_m1.paragraph_format.line_spacing = 1.3
    r = p_m1.add_run("A. Mô hình toán học cho chức năng Đăng ký tài khoản:\n")
    r.font.bold = True
    r.font.color.rgb = RGBColor(26, 82, 118)
    p_m1.add_run(
        "• Không gian đầu vào: X = { (email, password, confirm_password, full_name) }\n"
        "• Trạng thái cơ sở dữ liệu hiện hành: D = { u_1, u_2, ..., u_n } với u_i.email là các email đã tồn tại.\n"
        "• Điều kiện kiểm tra logic (Validation):\n"
        "  Valid(X) = (Format(email) == True) ∧ (Length(password) ≥ 8) ∧ (password == confirm_password) ∧ (∀u ∈ D: u.email ≠ email)\n"
        "• Hàm chuyển trạng thái (Mapping Function):\n"
        "  f_register(X) = { Thêm User(email, Hash(password), full_name) vào D → Thành công, Chuyển hướng /login/\n"
        "                  { Trả về thông báo lỗi tương ứng nếu Valid(X) == False"
    )

    p_m2 = doc.add_paragraph()
    p_m2.paragraph_format.left_indent = Inches(0.4)
    p_m2.paragraph_format.line_spacing = 1.3
    r = p_m2.add_run("B. Mô hình toán học cho chức năng Đăng nhập:\n")
    r.font.bold = True
    r.font.color.rgb = RGBColor(26, 82, 118)
    p_m2.add_run(
        "• Không gian đầu vào: Y = { (email, password) }\n"
        "• Hàm xác thực người dùng:\n"
        "  f_login(Y) = { Khởi tạo Session(u.id), Ghi Cookie, Phân luồng quyền nếu (∃u ∈ D: u.email == email ∧ u.is_active == True ∧ Verify(password, u.password_hash) == True)\n"
        "               { Từ chối truy cập và xuất thông báo lỗi nếu không thỏa mãn điều kiện xác thực."
    )

    doc.add_page_break()

    # ================= CHƯƠNG 2 =================
    h_c2 = doc.add_heading("CHƯƠNG 2: BIỂU ĐỒ VÀ THIẾT KẾ HỆ THỐNG", level=1)
    for r in h_c2.runs:
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(26, 82, 118)

    # 2.1
    h_21 = doc.add_heading("2.1. Biểu đồ Use Case", level=2)
    p = doc.add_paragraph(
        "Biểu đồ Use Case thể hiện tương tác giữa các tác nhân (Actors) bên ngoài với các chức năng bên trong hệ thống Urban Threads. "
        "Hệ thống bao gồm 3 tác nhân chính:\n"
        "1. Khách hàng (Customer / Guest): Đăng ký, đăng nhập, tìm kiếm sản phẩm, quản lý giỏ hàng, áp dụng coupon, quét mã VietQR thanh toán và xem lịch sử đơn hàng.\n"
        "2. Quản trị viên (Admin): Đăng nhập quản trị, quản lý danh mục và sản phẩm, quản lý biến thể tồn kho SKU, quản lý trạng thái đơn hàng, duyệt thanh toán 1-click qua ảnh biên lai.\n"
        "3. Cổng VietQR (External Service): Dịch vụ sinh mã VietQR động theo chuẩn Napas247 phục vụ giao dịch chuyển khoản."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    # Embed usecase bieu do.png if exists
    uc_img_path = r"D:\urban-threads\usercase bieu do.png"
    if os.path.exists(uc_img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(uc_img_path, width=Inches(5.5))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_c = p_cap.add_run("Hình 2.1: Biểu đồ Use Case tổng quan hệ thống Urban Threads")
        r_c.font.italic = True
        r_c.font.size = Pt(11)

    # 2.2
    h_22 = doc.add_heading("2.2. Biểu đồ tuần tự (Sequence Diagram)", level=2)
    p = doc.add_paragraph(
        "Biểu đồ tuần tự mô tả dòng thông điệp trao đổi giữa đối tượng Người dùng, Giao diện Web (Template/View), "
        "Bộ điều khiển (Controller/Middleware) và Cơ sở dữ liệu theo trình tự thời gian cho các kịch bản then chốt:"
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    # Sequence Dang ky
    h_221 = doc.add_heading("2.2.1. Biểu đồ tuần tự chức năng Đăng ký", level=3)
    p_seq1 = doc.add_paragraph(
        "1. Người dùng gửi yêu cầu đăng ký tại giao diện /accounts/register/.\n"
        "2. Giao diện gửi dữ liệu form (POST) đến Accounts View.\n"
        "3. View thực hiện validate dữ liệu và truy vấn Database để kiểm tra email trùng lặp.\n"
        "4. Nếu email chưa tồn tại: View gọi hàm băm PBKDF2 mã hóa mật khẩu và ghi User mới vào Database.\n"
        "5. Database xác nhận ghi thành công -> View gửi thông báo thành công và chuyển hướng người dùng sang trang Đăng nhập."
    )
    p_seq1.paragraph_format.line_spacing = 1.3
    p_seq1.paragraph_format.left_indent = Inches(0.4)

    # Sequence Dang nhap
    h_222 = doc.add_heading("2.2.2. Biểu đồ tuần tự chức năng Đăng nhập", level=3)
    p_seq2 = doc.add_paragraph(
        "1. Người dùng nhập Email và Mật khẩu tại giao diện /accounts/login/.\n"
        "2. Form gửi yêu cầu POST đến Accounts Login View.\n"
        "3. View truy vấn Database tìm người dùng tương ứng với Email.\n"
        "4. Nếu tìm thấy: Hệ thống gọi hàm check_password() để so khớp mật khẩu người dùng với chuỗi băm lưu trữ.\n"
        "5. Nếu hợp lệ: Django khởi tạo Session ID, lưu vào Cookie trình duyệt và chuyển hướng về trang chủ hoặc bảng điều khiển Admin (nếu là nhân viên)."
    )
    p_seq2.paragraph_format.line_spacing = 1.3
    p_seq2.paragraph_format.left_indent = Inches(0.4)

    # 2.3
    h_23 = doc.add_heading("2.3. Sơ đồ cơ sở dữ liệu (ER Diagram)", level=2)
    p = doc.add_paragraph(
        "Cơ sở dữ liệu của hệ thống Urban Threads được thiết kế theo mô hình quan hệ chuẩn hóa 3NF, bao gồm các bảng thực thể cốt lõi sau:\n"
        "• USER: Lưu trữ thông tin tài khoản người dùng (id, email, password, full_name, phone, is_staff, is_active).\n"
        "• CATEGORY: Danh mục sản phẩm thời trang (id, name, slug, description).\n"
        "• PRODUCT: Thông tin chung sản phẩm (id, category_id, name, slug, base_price, gender, is_active, is_featured).\n"
        "• PRODUCT_VARIANT: Biến thể chi tiết quản lý theo từng kích cỡ và màu sắc (id, product_id, color_id, size_id, sku, stock).\n"
        "• CART & CART_ITEM: Quản lý giỏ hàng lưu trữ session hoặc gắn với tài khoản người dùng.\n"
        "• ORDER & ORDER_ITEM: Quản lý đơn hàng (mã đơn DHxxxxxxxx, tổng tiền, địa chỉ giao nhận, trạng thái 7 bước) và chi tiết từng mặt hàng.\n"
        "• PAYMENT: Lưu trữ thông tin thanh toán chuyển khoản, mã ngân hàng bank_id, trạng thái và ảnh chụp biên lai upload của khách hàng."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    doc.add_page_break()

    # ================= CHƯƠNG 3 =================
    h_c3 = doc.add_heading("CHƯƠNG 3: TRÌNH BÀY DEMO HỆ THỐNG", level=1)
    for r in h_c3.runs:
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(26, 82, 118)

    demo_sections = [
        ("3.1. Giao diện Trang chủ và Danh mục", 
         "Trang chủ website Urban Threads (http://127.0.0.1:8000/) được thiết kế nổi bật với Hero Banner thời trang đường phố, "
         "khu vực hiển thị sản phẩm mới (New Arrivals), sản phẩm bán chạy (Best Sellers) và danh mục phân loại (Áo thun, Hoodie, Áo khoác, Quần jeans...). "
         "Các khối sản phẩm hiển thị đầy đủ hình ảnh chất lượng cao, giá bán định dạng chuẩn Việt Nam (200,000đ), nhãn trạng thái và nút thêm giỏ nhanh."),
        ("3.2. Giao diện Đăng nhập, Đăng ký và Tìm kiếm thông minh", 
         "Giao diện đăng ký và đăng nhập được thiết kế gọn gàng, hỗ trợ kiểm tra định dạng email và mật khẩu trực tiếp. "
         "Thanh tìm kiếm trên thanh điều hướng tích hợp thuật toán tìm kiếm tiếng Việt thông minh 2 tầng: hỗ trợ tìm kiếm cả từ khóa có dấu và không dấu "
         "(ví dụ: khách gõ \"ao khoac\" hay \"áo khoác\" đều trả về kết quả chính xác tuyệt đối)."),
        ("3.3. Trang chi tiết sản phẩm và chọn biến thể (Size / Màu)", 
         "Trang chi tiết sản phẩm cho phép người dùng xem thư viện ảnh gallery nhiều góc độ, xem thông tin chi tiết về chất liệu và hướng dẫn bảo quản. "
         "Người dùng có thể tương tác chọn màu sắc và kích thước (S, M, L, XL, XXL) để hệ thống tự động cập nhật số lượng tồn kho thực tế và mã SKU tương ứng."),
        ("3.4. Giỏ hàng và Ngăn kéo giỏ hàng (Mini Cart Drawer)", 
         "Khi bấm thêm vào giỏ, ngăn kéo Mini Cart sẽ trượt ra mượt mà từ cạnh phải màn hình mà không cần tải lại trang. "
         "Khách hàng có thể tăng/giảm số lượng sản phẩm, xóa sản phẩm, nhập mã giảm giá (Coupon như WELCOME10, GIAM50K) "
         "và quan sát thanh tiến trình đạt ngưỡng miễn phí ship (Freeship từ 500.000đ)."),
        ("3.5. Trang thanh toán và Cổng quét mã VietQR tự động", 
         "Tại trang thanh toán (Checkout), hệ thống hỗ trợ chọn địa chỉ từ sổ địa chỉ đã lưu và chọn phương thức vận chuyển. "
         "Đặc biệt, hệ thống sinh ra mã QR Code VietQR động chứa chính xác số tài khoản ngân hàng, tên thụ hưởng, số tiền thanh toán "
         "và cú pháp nội dung chuyển khoản DHxxxxxxxx. Khách hàng quét mã qua ứng dụng ngân hàng và tải lên ảnh chụp biên lai giao dịch thành công."),
        ("3.6. Trang quản trị Admin Dashboard", 
         "Trang quản trị Django Admin tùy biến cao cấp (http://127.0.0.1:8000/admin/) cung cấp đầy đủ công cụ kiểm soát: "
         "xem danh sách đơn hàng có gắn huy hiệu màu sắc trực quan (Status Badges), xem trước hình thu nhỏ ảnh biên lai của khách, "
         "bấm phóng to ảnh biên lai và thao tác 1-click để duyệt thanh toán, tự động chuyển đơn sang trạng thái Đang đóng gói."),
        ("3.7. Kế hoạch triển khai dự án (Gantt Chart)", 
         "Dự án được thực hiện nghiêm túc theo quy trình phát triển phần mềm gồm các giai đoạn: Khảo sát và phân tích yêu cầu (Tuần 1-2), "
         "Thiết kế kiến trúc hệ thống và cơ sở dữ liệu (Tuần 3), Lập trình module nghiệp vụ và giao diện người dùng (Tuần 4-5), "
         "Tích hợp cổng VietQR và kiểm thử hệ thống (Tuần 6), Viết tài liệu báo cáo và nghiệm thu đề tài.")
    ]

    for title_d, desc_d in demo_sections:
        h_d = doc.add_heading(title_d, level=2)
        for r in h_d.runs:
            r.font.bold = True
            r.font.size = Pt(13)
            r.font.color.rgb = RGBColor(26, 82, 118)
        p_d = doc.add_paragraph(desc_d)
        p_d.paragraph_format.line_spacing = 1.3
        p_d.paragraph_format.first_line_indent = Inches(0.4)

    doc.add_page_break()

    # ================= KẾT LUẬN =================
    h_kl = doc.add_heading("KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", level=1)
    for r in h_kl.runs:
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(26, 82, 118)

    p = doc.add_paragraph("1. Tóm tắt kết quả đạt được:")
    p.runs[0].font.bold = True
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "Sau quá trình nghiên cứu và thực hiện đề tài, em đã hoàn thành việc xây dựng website thương mại điện tử Urban Threads "
        "đáp ứng đầy đủ các yêu cầu đã đặt ra:\n"
        "• Hệ thống hoạt động ổn định trên nền tảng Django 5.x với cấu trúc module hóa rõ ràng (9 ứng dụng chuyên biệt).\n"
        "• Giao diện người dùng hiện đại, tinh tế theo phong cách streetwear, tương thích tốt trên cả thiết bị di động và máy tính.\n"
        "• Đầy đủ các tính năng thương mại điện tử cốt lõi: Quản lý tài khoản, duyệt danh mục, bộ lọc, tìm kiếm tiếng Việt thông minh, "
        "giỏ hàng phản ứng nhanh Mini Cart, mã giảm giá, tính phí ship linh hoạt.\n"
        "• Triển khai thành công giải pháp thanh toán quét mã VietQR tự động và cơ chế duyệt thanh toán 1-click đối soát biên lai trong Admin."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph("2. Đánh giá ưu điểm và hạn chế:")
    p.runs[0].font.bold = True
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "• Ưu điểm: Mã nguồn sạch, dễ bảo trì và mở rộng; cơ chế bảo mật cao nhờ các tính năng tích hợp sẵn của Django; "
        "trải nghiệm người dùng mượt mà, tiện lợi; hệ thống quản trị trực quan giúp người bán dễ dàng kiểm soát kinh doanh.\n"
        "• Hạn chế: Cơ chế đối soát thanh toán hiện vẫn cần bước nhân viên xác nhận hình ảnh biên lai (chưa tích hợp Webhook biến động số dư tự động của ngân hàng); "
        "chưa tích hợp tính năng trò chuyện trực tuyến (Live Chat) chăm sóc khách hàng tức thời."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph("3. Hướng phát triển trong tương lai:")
    p.runs[0].font.bold = True
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "• Tích hợp API ngân hàng hoặc Open Banking Webhook để tự động gạch nợ đơn hàng ngay khi khách hàng quét mã VietQR thành công.\n"
        "• Phát triển ứng dụng di động (Mobile App) trên Flutter hoặc React Native kết nối đồng bộ qua Django REST Framework.\n"
        "• Ứng dụng trí tuệ nhân tạo (AI Recommendation) để gợi ý các sản phẩm trang phục phối đồ (Mix & Match) cá nhân hóa cho từng khách hàng."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    doc.add_page_break()

    # ================= LỜI CẢM ƠN =================
    h_cm = doc.add_heading("LỜI CẢM ƠN", level=1)
    h_cm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in h_cm.runs:
        r.font.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor(26, 82, 118)

    p = doc.add_paragraph(
        "Em xin bày tỏ lòng biết ơn sâu sắc và chân thành nhất tới cô Phạm Thị Dung — giảng viên phụ trách học phần Công nghệ phần mềm. "
        "Trong suốt quá trình học tập và thực hiện đề tài, cô đã luôn tận tình hướng dẫn, chỉ bảo những kiến thức chuyên môn quý báu "
        "và định hướng phương pháp tư duy kỹ thuật phần mềm chuẩn mực."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "Những lời góp ý chi tiết, sự động viên kịp thời và phương pháp giảng dạy tâm huyết của cô đã giúp em vượt qua những bỡ ngỡ "
        "trong giai đoạn khảo sát yêu cầu, thiết kế kiến trúc và giải quyết các bài toán kỹ thuật phức tạp trong đề tài. "
        "Đây không chỉ là hành trang quan trọng để em hoàn thành tốt học phần mà còn là nền tảng vững chắc cho con đường phát triển nghề nghiệp sau này."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "Dù em đã rất nỗ lực cố gắng, song do kiến thức và kinh nghiệm thực tế của bản thân còn hạn chế, báo cáo khó tránh khỏi những thiếu sót nhất định. "
        "Em rất kính mong tiếp tục nhận được sự nhận xét, chỉ dẫn của cô để bài làm được hoàn thiện hơn nữa."
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p = doc.add_paragraph(
        "Em xin kính chúc cô luôn dồi dào sức khỏe, ngập tràn niềm vui, hạnh phúc và gặt hái được nhiều thành công hơn nữa trong sự nghiệp trồng người cao quý!\n\n"
    )
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.first_line_indent = Inches(0.4)

    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.paragraph_format.right_indent = Inches(0.5)
    r_s1 = p_sign.add_run("Sinh viên thực hiện\n\n\n\n")
    r_s1.font.italic = True
    r_s2 = p_sign.add_run("Lê Hải Nam")
    r_s2.font.bold = True
    r_s2.font.size = Pt(14)

    output_path = r"D:\urban-threads\Bao_Cao_Cong_Nghe_Phan_Mem_Le_Hai_Nam.docx"
    doc.save(output_path)
    print(f"Report saved successfully to: {output_path}")

if __name__ == "__main__":
    create_report()
