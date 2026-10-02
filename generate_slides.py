import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Colors
    COLOR_ORANGE_HEADER = RGBColor(217, 83, 30)      # #D9531E
    COLOR_ORANGE_SIDEBAR = RGBColor(217, 83, 30)
    COLOR_BLUE_TITLE = RGBColor(26, 82, 118)          # #1A5276
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_BLACK = RGBColor(20, 20, 20)
    COLOR_LIGHT_GRAY = RGBColor(245, 245, 245)
    COLOR_BORDER_GRAY = RGBColor(200, 200, 200)

    slides_data = [
        {
            "sub_title": "3.1. Giai đoạn lập kế hoạch – Khảo sát yêu cầu (Tài khoản & Người dùng)",
            "side_text": "Viết\nyêu\ncầu\nmột\ncách\nrõ\nràng",
            "rows": [
                ["Mã", "Tên yêu cầu", "Mô tả Yêu cầu", "Ưu tiên", "Tiêu chí nghiệm thu", "Tình trạng thực hiện", "Link mẫu"],
                ["FR-01", "Đăng ký tài khoản", "Người dùng có thể đăng ký tài khoản mới bằng Email và mật khẩu bảo mật.", "Cao", "Nhập đúng thông tin tạo tài khoản thành công; trùng email hoặc sai định dạng hiển thị thông báo lỗi.", "Đã hoàn thành", "chivalry-warehouse.com"],
                ["FR-02", "Đăng nhập / Đăng xuất", "Người dùng đăng nhập bằng Email/Password, có ghi nhớ đăng nhập và đăng xuất an toàn.", "Cao", "Đăng nhập đúng vào trang cá nhân; sai mật khẩu báo lỗi; đăng xuất hủy session an toàn.", "Đã hoàn thành", "representclo.com"],
                ["FR-03", "Sổ địa chỉ giao hàng", "Người dùng quản lý nhiều địa chỉ nhận hàng (thêm, sửa, xóa, đặt địa chỉ mặc định).", "Trung bình", "Dữ liệu địa chỉ lưu vào tài khoản; tự động điền địa chỉ mặc định khi đặt hàng.", "Đã hoàn thành", "culturekings.com.au"],
                ["FR-04", "Lịch sử đơn hàng", "Người dùng theo dõi danh sách đơn hàng đã mua và tiến trình trạng thái chi tiết.", "Cao", "Hiển thị mã đơn DHxxxx, danh sách sản phẩm, giá tiền, ngày đặt và timeline trạng thái.", "Đã hoàn thành", "representclo.com"]
            ]
        },
        {
            "sub_title": "3.1. Giai đoạn lập kế hoạch – Khảo sát yêu cầu (Sản phẩm & Tìm kiếm)",
            "side_text": "Viết\nyêu\ncầu\nmột\ncách\nrõ\nràng",
            "rows": [
                ["Mã", "Tên yêu cầu", "Mô tả Yêu cầu", "Ưu tiên", "Tiêu chí nghiệm thu", "Tình trạng thực hiện", "Link mẫu"],
                ["FR-05", "Duyệt danh mục & Bộ lọc", "Phân loại sản phẩm (Áo, Quần, Phụ kiện, Sale) kèm bộ lọc theo giá, màu, size.", "Cao", "Lọc chính xác theo nhiều tiêu chí đồng thời; sắp xếp theo giá tăng/giảm, mới nhất.", "Đã hoàn thành", "chivalry-warehouse.com"],
                ["FR-06", "Tìm kiếm thông minh", "Tìm kiếm sản phẩm theo từ khóa tiếng Việt đa tầng (hỗ trợ cả có dấu và không dấu).", "Cao", "Gõ 'ao thun' hay 'áo thun' đều trả về kết quả đúng; có popup gợi ý sản phẩm.", "Đã hoàn thành", "culturekings.com.au"],
                ["FR-07", "Chi tiết & Biến thể SKU", "Trang chi tiết sản phẩm hiển thị gallery ảnh, bảng chọn size/màu và tồn kho tương ứng.", "Cao", "Đổi màu/size tự động cập nhật mã SKU, giá bán và cảnh báo còn/hết hàng tức thì.", "Đã hoàn thành", "representclo.com"],
                ["FR-08", "Danh sách yêu thích", "Khách hàng có thể lưu các sản phẩm yêu thích (Wishlist) vào tài khoản cá nhân.", "Trung bình", "Bấm icon trái tim thêm/xóa nhanh bằng AJAX mà không cần tải lại trang.", "Đã hoàn thành", "chivalry-warehouse.com"]
            ]
        },
        {
            "sub_title": "3.1. Giai đoạn lập kế hoạch – Khảo sát yêu cầu (Giỏ hàng & Thanh toán)",
            "side_text": "Viết\nyêu\ncầu\nmột\ncách\nrõ\nràng",
            "rows": [
                ["Mã", "Tên yêu cầu", "Mô tả Yêu cầu", "Ưu tiên", "Tiêu chí nghiệm thu", "Tình trạng thực hiện", "Link mẫu"],
                ["FR-09", "Giỏ hàng & Mini Cart Drawer", "Thêm sản phẩm vào giỏ, xem nhanh qua ngăn kéo trượt ra từ bên phải màn hình.", "Cao", "Mini cart trượt mượt mà; tăng giảm số lượng, xóa item và cập nhật tổng tiền realtime.", "Đã hoàn thành", "culturekings.com.au"],
                ["FR-10", "Mã giảm giá & Freeship bar", "Áp dụng coupon khuyến mãi và hiển thị thanh tiến trình đạt ngưỡng miễn phí ship.", "Cao", "Nhập mã hợp lệ tự trừ tiền; đơn hàng đạt từ 500,000đ tự động miễn phí vận chuyển.", "Đã hoàn thành", "culturekings.com.au"],
                ["FR-11", "Thanh toán quét mã VietQR", "Tự động sinh mã VietQR động theo chuẩn Napas247 cho bất kỳ ngân hàng nào.", "Cao", "Mã QR chứa đúng số tiền, STK và nội dung DHxxxx; có nút 1-click copy nhanh.", "Đã hoàn thành", "VietQR / chivalry-warehouse.com"],
                ["FR-12", "Upload biên lai đối soát", "Khách hàng tải lên ảnh chụp màn hình chuyển khoản sau khi quét mã QR.", "Cao", "Tải được ảnh JPG/PNG; tự động cập nhật trạng thái đơn sang 'Đã xác nhận'.", "Đã hoàn thành", "Urban Threads Custom"]
            ]
        },
        {
            "sub_title": "3.1. Giai đoạn lập kế hoạch – Khảo sát yêu cầu (Quản trị Admin Dashboard)",
            "side_text": "Viết\nyêu\ncầu\nmột\ncách\nrõ\nràng",
            "rows": [
                ["Mã", "Tên yêu cầu", "Mô tả Yêu cầu", "Ưu tiên", "Tiêu chí nghiệm thu", "Tình trạng thực hiện", "Link mẫu"],
                ["FR-13", "Quản lý sản phẩm & Kho", "Admin thêm, sửa, xóa sản phẩm, danh mục, gallery ảnh và quản lý tồn kho biến thể.", "Cao", "Dữ liệu lưu chuẩn CSDL; có thumbnail ảnh trực quan; kiểm soát tồn kho từng size/màu.", "Đã hoàn thành", "Django Admin / representclo.com"],
                ["FR-14", "Quản lý đơn hàng", "Theo dõi vòng đời đơn hàng qua 7 trạng thái, lọc theo thời gian và bulk update.", "Cao", "Cập nhật trạng thái chuẩn xác; tự động trừ tồn kho khi đơn hàng được xác nhận.", "Đã hoàn thành", "Django Admin"],
                ["FR-15", "Duyệt thanh toán 1-click", "Admin kiểm tra ảnh biên lai do khách gửi và duyệt thanh toán trực tiếp tại bảng đơn.", "Cao", "Xem trước ảnh bill phóng to; bấm 1 nút xác nhận tiền về và chuyển đơn sang đóng gói.", "Đã hoàn thành", "Urban Threads Custom"],
                ["FR-16", "Quản lý khuyến mãi & Banner", "Tạo mã coupon (% hoặc số tiền), giới hạn số lượt dùng và quản lý banner trang chủ.", "Trung bình", "Mã hết hạn hoặc hết lượt sẽ tự động từ chối; banner cập nhật đúng thứ tự hiển thị.", "Đã hoàn thành", "culturekings.com.au"]
            ]
        },
        {
            "sub_title": "3.1. Giai đoạn lập kế hoạch – Khảo sát yêu cầu (Yêu cầu Phi chức năng - NFR)",
            "side_text": "Viết\nyêu\ncầu\nmột\ncách\nrõ\nràng",
            "rows": [
                ["Mã", "Tên yêu cầu", "Mô tả Yêu cầu", "Ưu tiên", "Tiêu chí nghiệm thu", "Tình trạng thực hiện", "Link mẫu"],
                ["NFR-01", "Hiệu năng tải trang", "Trang chủ, danh mục và chi tiết sản phẩm phản hồi nhanh trong thời gian chấp nhận được.", "Cao", "Thời gian phản hồi trang dưới 2 giây với kết nối thông thường; tối ưu query ORM.", "Đã hoàn thành", "Google Lighthouse / chivalry-warehouse.com"],
                ["NFR-02", "Bảo mật hệ thống", "Bảo vệ thông tin người dùng, an toàn giao dịch và ngăn chặn các lỗ hổng phổ biến.", "Cao", "Bảo vệ CSRF token cho mọi form, chống XSS, SQL Injection và mã hóa mật khẩu.", "Đã hoàn thành", "Django Security Core"],
                ["NFR-03", "Thiết kế Responsive", "Giao diện hiển thị tối ưu trên mọi thiết bị: Điện thoại, Máy tính bảng và Laptop.", "Cao", "Layout co giãn tự động theo Tailwind CSS; menu mobile và cart drawer mượt mà.", "Đã hoàn thành", "representclo.com"],
                ["NFR-04", "Linh hoạt cấu hình QR Bank", "Hệ thống hỗ trợ cấu hình chuyển đổi sang bất kỳ tài khoản ngân hàng nào tại Việt Nam.", "Cao", "Đổi BANK_ID trong .env (VCB, TCB, MB, ACB...) hệ thống tự nhận diện tên ngân hàng.", "Đã hoàn thành", "VietQR Standard API"]
            ]
        }
    ]

    blank_slide_layout = prs.slide_layouts[6]

    for page_idx, sdata in enumerate(slides_data):
        slide = prs.slides.add_slide(blank_slide_layout)

        # 1. Header bar
        header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.333), Inches(0.8))
        tf = header_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = sdata["sub_title"]
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = COLOR_BLUE_TITLE
        p.font.name = "Arial"

        # 2. Left side banner (Orange box)
        left_box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.5), Inches(1.3), Inches(2.0), Inches(5.7)
        )
        left_box.fill.solid()
        left_box.fill.fore_color.rgb = COLOR_ORANGE_SIDEBAR
        left_box.line.color.rgb = COLOR_ORANGE_SIDEBAR
        
        ltf = left_box.text_frame
        ltf.word_wrap = True
        lp = ltf.paragraphs[0]
        lp.text = sdata["side_text"]
        lp.font.size = Pt(28)
        lp.font.bold = True
        lp.font.color.rgb = COLOR_WHITE
        lp.font.name = "Arial"
        lp.alignment = PP_ALIGN.CENTER

        # 3. Table
        table_rows = len(sdata["rows"])
        table_cols = 7
        table_shape = slide.shapes.add_table(
            table_rows, table_cols,
            Inches(2.7), Inches(1.3), Inches(10.1), Inches(5.7)
        )
        table = table_shape.table

        # Column widths
        # Total = 10.1 inches
        col_widths = [
            Inches(0.9),   # Mã
            Inches(1.5),   # Tên yêu cầu
            Inches(2.3),   # Mô tả Yêu cầu
            Inches(0.9),   # Ưu tiên
            Inches(2.4),   # Tiêu chí nghiệm thu
            Inches(1.1),   # Tình trạng
            Inches(1.0)    # Link mẫu
        ]
        for c_idx, width in enumerate(col_widths):
            table.columns[c_idx].width = width

        # Populate cells
        for r_idx, row in enumerate(sdata["rows"]):
            for c_idx, val in enumerate(row):
                cell = table.cell(r_idx, c_idx)
                cell.text = val
                cp = cell.text_frame.paragraphs[0]
                cp.font.name = "Arial"

                if r_idx == 0:
                    # Header row
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = COLOR_ORANGE_HEADER
                    cp.font.bold = True
                    cp.font.size = Pt(13)
                    cp.font.color.rgb = COLOR_WHITE
                    cp.alignment = PP_ALIGN.CENTER
                else:
                    # Data row
                    cp.font.size = Pt(11)
                    cp.font.color.rgb = COLOR_BLACK
                    if c_idx in [0, 3, 5, 6]:
                        cp.alignment = PP_ALIGN.CENTER
                    else:
                        cp.alignment = PP_ALIGN.LEFT
                    
                    if r_idx % 2 == 1:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = COLOR_WHITE
                    else:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = COLOR_LIGHT_GRAY

    output_path = r"D:\urban-threads\slide_khao_sat_yeu_cau_urban_threads.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_deck()
